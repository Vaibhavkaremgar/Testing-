from __future__ import annotations

import tempfile
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Any, Dict, List, Optional, Tuple
from unittest.mock import patch

import fitz
from docx import Document
from PIL import Image, ImageDraw

from ats.extraction.resume_parser import (
    _extract_pdf_text_via_ocr,
    _extract_pdf_text_with_pypdf,
    _has_meaningful_text,
    extract_document,
    parse_resume_text,
)


FIELD_NAMES = ["name", "email", "phone", "location", "current_role", "current_company"]


@dataclass
class FormatBenchmarkCase:
    case_id: str
    label: str
    filename: str
    expected: Dict[str, Any]
    mocked_ocr_text: Optional[str] = None


def _insert_lines(
    page: fitz.Page,
    lines: List[str],
    x: float,
    y: float,
    width: float,
    color: Tuple[float, float, float] = (0, 0, 0),
    fontsize: float = 11,
) -> None:
    cursor_y = y
    for line in lines:
        page.insert_text((x, cursor_y), line, fontsize=fontsize, fontname="helv", color=color)
        cursor_y += 16


def _create_multi_column_resume(path: Path) -> FormatBenchmarkCase:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    _insert_lines(
        page,
        [
            "Ananya Krishnan",
            "Hyderabad, Telangana | +91 98001 23456 | ananya.krishnan@outlook.com",
        ],
        40,
        40,
        500,
        fontsize=12,
    )
    _insert_lines(page, ["TECHNICAL SKILLS", "SQL", "Python", "Power BI", "Tableau", "Mixpanel", "Git"], 320, 110, 190)
    _insert_lines(
        page,
        [
            "WORK EXPERIENCE",
            "Senior Data Analyst",
            "Meesho Pvt. Ltd.",
            "Jun 2022 - Present",
            "Built reporting dashboards and business insights.",
        ],
        40,
        110,
        220,
    )
    doc.save(path)
    doc.close()
    return FormatBenchmarkCase(
        case_id="multi_column",
        label="Multi-column PDF",
        filename=path.name,
        expected={
            "name": "Ananya Krishnan",
            "email": "ananya.krishnan@outlook.com",
            "phone": "+91 98001 23456",
            "location": "Hyderabad, Telangana",
            "current_role": "Senior Data Analyst",
            "current_company": "Meesho Pvt. Ltd",
        },
    )


def _create_colored_resume(path: Path) -> FormatBenchmarkCase:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(0, 0, 595, 110), color=(0.12, 0.36, 0.64), fill=(0.12, 0.36, 0.64))
    _insert_lines(
        page,
        ["Priya Menon", "Bengaluru, Karnataka | +91 99887 66554 | priya.menon@email.com"],
        40,
        36,
        500,
        color=(1, 1, 1),
        fontsize=12,
    )
    _insert_lines(
        page,
        [
            "WORK EXPERIENCE",
            "Senior Product Designer",
            "Nova Design Studio",
            "2021 - Present",
            "Led design systems and UX research for hiring tools.",
        ],
        40,
        140,
        500,
    )
    doc.save(path)
    doc.close()
    return FormatBenchmarkCase(
        case_id="colored",
        label="Colored PDF",
        filename=path.name,
        expected={
            "name": "Priya Menon",
            "email": "priya.menon@email.com",
            "phone": "+91 99887 66554",
            "location": "Bengaluru, Karnataka",
            "current_role": "Senior Product Designer",
            "current_company": "Nova Design Studio",
        },
    )


def _create_horizontal_resume(path: Path) -> FormatBenchmarkCase:
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    _insert_lines(
        page,
        [
            "Maya Thomas",
            "Austin, Texas | +1 (415) 555-0101 | maya.thomas.engineer@gmail.com",
            "WORK EXPERIENCE",
            "Principal Backend Engineer",
            "Acme Cloud Systems",
            "2023 - Present",
        ],
        40,
        50,
        330,
    )
    _insert_lines(page, ["SKILLS", "Python | FastAPI | AWS | Docker | PostgreSQL"], 430, 50, 320)
    doc.save(path)
    doc.close()
    return FormatBenchmarkCase(
        case_id="horizontal",
        label="Horizontal PDF",
        filename=path.name,
        expected={
            "name": "Maya Thomas",
            "email": "maya.thomas.engineer@gmail.com",
            "phone": "+1 (415) 555-0101",
            "location": "Austin, Texas",
            "current_role": "Principal Backend Engineer",
            "current_company": "Acme Cloud Systems",
        },
    )


def _create_table_resume(path: Path) -> FormatBenchmarkCase:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    _insert_lines(page, ["Rahul Verma", "Pune, Maharashtra | +91 9876543210 | rahul.verma.ai.dev@gmail.com"], 40, 40, 500, fontsize=12)
    left = 40
    top = 120
    row_height = 28
    columns = [left, 220, 400, 555]
    for row in range(6):
        y = top + row * row_height
        page.draw_line((left, y), (555, y), color=(0, 0, 0), width=0.6)
    page.draw_line((left, top + 6 * row_height), (555, top + 6 * row_height), color=(0, 0, 0), width=0.6)
    for x in columns:
        page.draw_line((x, top), (x, top + 6 * row_height), color=(0, 0, 0), width=0.6)
    rows = [
        ("Section", "Title", "Company / Value"),
        ("Experience", "Senior ML Engineer", "AI Labs Pvt Ltd"),
        ("Dates", "2022 - Present", "Bangalore"),
        ("Skills", "Python | ML | NLP", "AWS | Docker"),
        ("Education", "B.Tech", "NIT Trichy"),
        ("Email", "rahul.verma.ai.dev@gmail.com", "+91 9876543210"),
    ]
    for index, row in enumerate(rows):
        y = top + index * row_height + 8
        for col_index, value in enumerate(row):
            page.insert_text((columns[col_index] + 4, y), value, fontsize=10, fontname="helv")
    doc.save(path)
    doc.close()
    return FormatBenchmarkCase(
        case_id="table_based",
        label="Table-based PDF",
        filename=path.name,
        expected={
            "name": "Rahul Verma",
            "email": "rahul.verma.ai.dev@gmail.com",
            "phone": "+91 9876543210",
            "location": "Pune, Maharashtra",
            "current_role": "Senior ML Engineer",
            "current_company": "AI Labs Pvt Ltd",
        },
    )


def _create_docx_resume(path: Path) -> FormatBenchmarkCase:
    document = Document()
    document.add_paragraph("Nisha Verma")
    document.add_paragraph("Pune, Maharashtra | +91 98765 11111 | nisha.verma@email.com")
    document.add_paragraph("Work Experience")
    document.add_paragraph("Talent Acquisition Specialist")
    document.add_paragraph("Bright Hire Solutions")
    document.add_paragraph("2022 - Present")
    skills_table = document.add_table(rows=2, cols=2)
    skills_table.cell(0, 0).text = "Skills"
    skills_table.cell(0, 1).text = "Sourcing | Screening"
    skills_table.cell(1, 0).text = "Tools"
    skills_table.cell(1, 1).text = "LinkedIn Recruiter | Greenhouse"
    document.save(path)
    return FormatBenchmarkCase(
        case_id="docx",
        label="DOCX resume",
        filename=path.name,
        expected={
            "name": "Nisha Verma",
            "email": "nisha.verma@email.com",
            "phone": "+91 98765 11111",
            "location": "Pune, Maharashtra",
            "current_role": "Talent Acquisition Specialist",
            "current_company": "Bright Hire Solutions",
        },
    )


def _create_image_resume(path: Path) -> FormatBenchmarkCase:
    image = Image.new("RGB", (1200, 1600), color=(255, 255, 255))
    drawer = ImageDraw.Draw(image)
    drawer.text((40, 40), "Ravi Kumar", fill=(0, 0, 0))
    drawer.text((40, 90), "Hyderabad, Telangana", fill=(0, 0, 0))
    image.save(path)
    ocr_text = "\n".join(
        [
            "Ravi Kumar",
            "Hyderabad, Telangana | +91 99887 66554 | ravi.kumar@email.com",
            "WORK EXPERIENCE",
            "Backend Engineer",
            "Acme Systems Ltd",
            "Jan 2022 - Present",
        ]
    )
    return FormatBenchmarkCase(
        case_id="image_converted",
        label="Image-converted resume",
        filename=path.name,
        expected={
            "name": "Ravi Kumar",
            "email": "ravi.kumar@email.com",
            "phone": "+91 99887 66554",
            "location": "Hyderabad, Telangana",
            "current_role": "Backend Engineer",
            "current_company": "Acme Systems Ltd",
        },
        mocked_ocr_text=ocr_text,
    )


def _build_sample_cases(root: Path) -> List[FormatBenchmarkCase]:
    return [
        _create_multi_column_resume(root / "multi_column_resume.pdf"),
        _create_colored_resume(root / "colored_resume.pdf"),
        _create_horizontal_resume(root / "horizontal_resume.pdf"),
        _create_table_resume(root / "table_resume.pdf"),
        _create_docx_resume(root / "docx_resume.docx"),
        _create_image_resume(root / "image_resume.png"),
    ]


def _legacy_extract_text(file_path: str) -> str:
    if file_path.lower().endswith(".pdf"):
        text_parts, _ = _extract_pdf_text_with_pypdf(file_path)
        if not _has_meaningful_text(text_parts):
            ocr_parts = _extract_pdf_text_via_ocr(file_path)
            if _has_meaningful_text(ocr_parts):
                text_parts = ocr_parts
        return "\n".join(part.strip() for part in text_parts if part and part.strip())
    return ""


def _field_accuracy(parsed: Dict[str, Any], expected: Dict[str, Any]) -> float:
    scores: List[float] = []
    for field in FIELD_NAMES:
        expected_value = " ".join(str(expected.get(field) or "").split()).lower()
        actual_value = " ".join(str(parsed.get(field) or "").split()).lower()
        if not expected_value:
            continue
        scores.append(1.0 if actual_value == expected_value else 0.0)
    return round(sum(scores) / len(scores), 3) if scores else 0.0


def _snippet(text: str, limit: int = 220) -> str:
    compact = " ".join((text or "").split())
    if len(compact) <= limit:
        return compact
    return compact[: limit - 3] + "..."


def _parse_case(case: FormatBenchmarkCase, file_path: Path) -> Dict[str, Any]:
    if case.mocked_ocr_text is None:
        document = extract_document(str(file_path))
    else:
        mock_parts = [case.mocked_ocr_text]
        if file_path.suffix.lower() == ".png":
            with patch("ats.extraction.resume_parser._extract_image_text_via_ocr", return_value=mock_parts):
                document = extract_document(str(file_path))
        else:
            with patch("ats.extraction.resume_parser._extract_pdf_text_via_ocr", return_value=mock_parts):
                document = extract_document(str(file_path))
    current_text = str(document.get("text") or "")
    current_layout = document.get("layout") or {}
    return {
        "document": document,
        "parsed": parse_resume_text(current_text, original_filename=case.filename, layout_signals=current_layout),
        "text": current_text,
        "layout": current_layout,
    }


def run_resume_format_benchmark() -> Dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="resume_format_benchmark_") as temp_dir:
        root = Path(temp_dir)
        cases = _build_sample_cases(root)
        case_results: List[Dict[str, Any]] = []
        before_scores: List[float] = []
        after_scores: List[float] = []

        for case in cases:
            file_path = root / case.filename
            legacy_text = _legacy_extract_text(str(file_path))
            parsed_case = _parse_case(case, file_path)
            current_parsed = parsed_case["parsed"]
            current_text = parsed_case["text"]
            current_layout = parsed_case["layout"]
            legacy_parsed = parse_resume_text(legacy_text, original_filename=case.filename) if legacy_text else {}

            before_score = _field_accuracy(legacy_parsed, case.expected) if legacy_text else 0.0
            after_score = _field_accuracy(current_parsed, case.expected)
            before_scores.append(before_score)
            after_scores.append(after_score)

            case_results.append(
                {
                    "id": case.case_id,
                    "label": case.label,
                    "expected": case.expected,
                    "before": {
                        "score": before_score,
                        "parsed": {field: legacy_parsed.get(field) for field in FIELD_NAMES},
                        "snippet": _snippet(legacy_text),
                    },
                    "after": {
                        "score": after_score,
                        "parsed": {field: current_parsed.get(field) for field in FIELD_NAMES},
                        "snippet": _snippet(current_text),
                        "layout_signals": current_layout,
                        "confidence": current_parsed.get("confidence", {}),
                    },
                    "improvement": round(after_score - before_score, 3),
                }
            )

        before_avg = round(mean(before_scores), 3) if before_scores else 0.0
        after_avg = round(mean(after_scores), 3) if after_scores else 0.0
        return {
            "total_cases": len(cases),
            "before_accuracy": before_avg,
            "after_accuracy": after_avg,
            "absolute_improvement": round(after_avg - before_avg, 3),
            "relative_improvement_percent": round(((after_avg - before_avg) / before_avg) * 100, 1) if before_avg else None,
            "case_results": case_results,
        }
