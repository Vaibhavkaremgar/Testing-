from __future__ import annotations

import logging
import os
import re
import shutil
import time
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Dict, Iterable, List, Optional, Tuple

import pdfplumber
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image
try:
    import fitz
except ImportError:  # pragma: no cover - optional dependency
    fitz = None

try:
    from docx2python import docx2python
except ImportError:  # pragma: no cover - optional dependency
    docx2python = None

try:
    import mammoth
except ImportError:  # pragma: no cover - optional dependency
    mammoth = None

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency
    pytesseract = None

from ats.datasets.parser_config_loader import ParserConfigLoader
from ats.extraction.entity_extraction import extract_resume_entities
from ats.extraction.experience_extraction import compute_total_experience, parse_date
from ats.extraction.information_extraction import extract_email as extract_normalized_email
from ats.extraction.information_extraction import extract_location as extract_normalized_location
from ats.extraction.information_extraction import extract_resume_information
from ats.extraction.layout_detection import get_layout_runtime_status, infer_layout_signals
from ats.extraction.postprocessing import apply_postprocessing, dedupe_strings
from ats.extraction.resume_type_detection import detect_resume_type
from ats.extraction.validation import validate_location, validate_parsed_fields
from ats.preprocessing.section_segmentation import segment_resume_sections
from ats.preprocessing.text_cleaning import clean_text, clean_text_pipeline, normalize_common_artifacts, normalize_document_structure, split_inline_section_headers
from app.spacy_nlp import SPACY_AVAILABLE, get_section_doc

logger = logging.getLogger(__name__)

PHONE_PATTERNS = [
    r"\+91[-\s]?\d{5}[-\s]?\d{5}",
    r"\+91[-\s]?\d{10}",
    r"\d{5}[-\s]?\d{5}",
    r"\+?1?[-\s]?\(?\d{3}\)?[-\s]?\d{3}[-\s]?\d{4}",
    r"\(?\d{3}\)?[-\s]?\d{3}[-\s]?\d{4}",
]
EMAIL_PATTERN = re.compile(r"\b[A-Za-z0-9._%+-]+\s*@\s*[A-Za-z0-9.-]+\s*\.\s*[A-Za-z]{2,}\b")
PHONE_LINE_PATTERN = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
NAME_LINE_DISALLOWED_PATTERN = re.compile(
    r"(?i)\b(?:resumev|resume|cv|profile|skills|education|experience|contact|objective|snapshot|summary|details|name)\b"
)
TEXTBOX_TEXT_PATTERN = re.compile(r"<w:t[^>]*>(.*?)</w:t>", re.IGNORECASE | re.DOTALL)
DEFAULT_INVALID_NAME_TOKENS = {
    "about", "machine", "learning", "python", "java", "react", "sql", "developer",
    "engineer", "manager", "analyst", "summary", "profile", "objective", "resume",
    "curriculum", "vitae", "experience", "skills", "education", "project", "projects",
    "email", "phone", "address", "location", "contact", "details", "snapshot", "professional", "job", "name",
}
DEFAULT_NAME_STOP_TOKENS = {
    "senior", "sr", "junior", "jr", "principal", "staff", "assistant",
    "frontend", "front-end", "backend", "back-end", "fullstack", "full-stack",
    "software", "data", "product", "business", "human", "resources", "hr",
    "cybersecurity", "security", "cloud", "automation", "test", "testing", "qa",
    "network", "soc", "penetration", "mobile", "web", "python", "java", "typescript",
    "playwright", "selenium", "robot", "framework", "analyst", "tester",
    "engineer", "developer", "manager", "analyst", "scientist", "consultant", "architect",
    "specialist", "designer", "executive", "director", "associate", "lead", "intern",
    "summary", "profile", "objective", "experience", "skills", "education", "projects",
    "location", "email", "phone", "mobile", "linkedin", "github",
}
HEADER_NAME_SPLIT_PATTERN = re.compile(r"\s+[|,/-]\s+|\s{2,}")
INLINE_CONTACT_PATTERN = re.compile(
    r"(?i)(\+?\d[\d\s().-]{7,}\d|[A-Za-z0-9._%+-]+\s*@\s*[A-Za-z0-9.-]+\s*\.\s*[A-Za-z]{2,}|linkedin|github|portfolio)"
)
NAME_LABEL_PATTERN = re.compile(r"(?i)^\s*name\s*[:\-]\s*(?P<value>.+)$")
SECTION_START_PATTERN = re.compile(
    r"(?i)^(?:work experience|professional experience|employment history|employment|career history|experience|period|"
    r"skills|technical skills|core skills|key skills|education|projects?|summary|profile|languages?|"
    r"certifications?|achievements?|awards?|publications?|references?|professional snapshot|snapshot|overview)$"
)
NAME_CONTEXT_ROLE_PATTERN = re.compile(
    r"(?i)\b(?:engineer|developer|tester|analyst|consultant|manager|architect|specialist|intern)\b"
)
NAME_COMPANY_PATTERN = re.compile(
    r"(?i)\b(?:pvt|ltd|inc|llc|llp|corp|corporation|technologies|technology|solutions|systems|labs|works|school|college|university|academy|institute|services)\b"
)
UPPERCASE_NAME_PATTERN = re.compile(r"^[A-Z][A-Z'`.-]*(?:\s+[A-Z][A-Z'`.-]*){1,3}$")
PDF_LINE_TOLERANCE = 3.0
PDF_MIN_COLUMN_GAP = 60.0
PDF_MIN_LINES_PER_COLUMN = 8
PDF_SEGMENT_GAP = 35.0
OCR_MIN_TEXT_LENGTH = 80
OCR_MIN_ALPHA_CHARS = 30
OCR_MIN_ALPHA_RATIO = 0.3
BROKEN_TOKEN_PATTERN = re.compile(r"\b[A-Za-z]{1,8}\s+[A-Za-z]{1,8}\b")
BROKEN_MONTH_PATTERN = re.compile(
    r"(?i)\b(?:j\s+anuary|f\s+ebruary|m\s+arch|a\s+pril|m\s+ay|j\s+une|j\s+uly|s\s+eptember|o\s+ctober|n\s+ovember|d\s+ecember)\b"
)
SPLIT_EMAIL_ARTIFACT_PATTERN = re.compile(r"(?i)\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\s+[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
PDF_PARSER_PREFERENCE = {"pymupdf": 2, "pdfplumber": 1}
CONFIDENCE_RETRY_THRESHOLD = 0.7

_parser_config_loader = ParserConfigLoader()
_parser_vocabulary = _parser_config_loader.load_parser_vocabulary()
INVALID_NAME_TOKENS = set(
    str(value).strip().lower()
    for value in (_parser_vocabulary.get("invalid_name_tokens") or DEFAULT_INVALID_NAME_TOKENS)
    if str(value).strip()
)
NAME_STOP_TOKENS = set(
    str(value).strip().lower()
    for value in (_parser_vocabulary.get("name_stop_tokens") or DEFAULT_NAME_STOP_TOKENS)
    if str(value).strip()
)


def _iter_docx_blocks(parent):
    parent_element = parent.element.body if isinstance(parent, DocxDocument) else parent._tc
    for child in parent_element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def _extract_docx_table_lines(table: Table) -> List[str]:
    lines: List[str] = []
    for row in table.rows:
        row_values: List[str] = []
        for cell in row.cells:
            cell_parts: List[str] = []
            for block in _iter_docx_blocks(cell):
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if text:
                        cell_parts.append(text)
                elif isinstance(block, Table):
                    nested_lines = _extract_docx_table_lines(block)
                    if nested_lines:
                        cell_parts.extend(nested_lines)
            cell_text = " ".join(part.strip() for part in cell_parts if part.strip()).strip()
            if cell_text:
                row_values.append(cell_text)
        if row_values:
            lines.append(" | ".join(row_values))
    return lines


def _normalize_docx_line(value: str) -> str:
    line = normalize_common_artifacts(str(value or ""))
    line = line.replace("\r", "\n")
    line = re.sub(r"[ \t]+", " ", line)
    return line.strip()


def _dedupe_preserve_order(lines: Iterable[str]) -> List[str]:
    ordered: List[str] = []
    seen: set[str] = set()
    for line in lines:
        normalized = _normalize_docx_line(line)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        ordered.append(normalized)
    return ordered


def _is_valid_name_line(line: str) -> bool:
    candidate = _normalize_docx_line(line)
    if not candidate:
        return False
    if NAME_LINE_DISALLOWED_PATTERN.search(candidate):
        return False
    if EMAIL_PATTERN.search(candidate) or PHONE_LINE_PATTERN.search(candidate):
        return False
    if any(char.isdigit() for char in candidate):
        return False
    words = [word for word in candidate.split() if word]
    if not (2 <= len(words) <= 4):
        return False
    for word in words:
        normalized_word = word.strip(".,")
        if not normalized_word.replace("-", "").replace("'", "").isalpha():
            return False
        if not normalized_word[:1].isupper():
            return False
    return True


def extract_name(text: str) -> str:
    """Extract the candidate name from the first 10-15 lines, skipping resume/file-name noise."""
    if not text:
        return ""

    lines = [line.strip() for line in clean_text(text).splitlines() if line.strip()]
    for line in lines[:15]:
        if _is_valid_name_line(line):
            return _normalize_docx_line(line)
    return ""


def _flatten_docx2python_node(node: Any, *, depth: int = 0) -> List[str]:
    if node is None:
        return []
    if isinstance(node, str):
        parts = [part.strip() for part in node.replace("\r", "\n").split("\n")]
        return [part for part in parts if part]
    if not isinstance(node, (list, tuple)):
        return []

    if node and all(isinstance(item, (list, tuple)) for item in node):
        child_rows = [_flatten_docx2python_node(item, depth=depth + 1) for item in node]
        if depth >= 2 and any(child_rows):
            row_values = [" ".join(value for value in row if value).strip() for row in child_rows]
            row_values = [value for value in row_values if value]
            if row_values:
                return [" | ".join(row_values)]

    flattened: List[str] = []
    for item in node:
        flattened.extend(_flatten_docx2python_node(item, depth=depth + 1))
    return flattened


def _extract_docx_headers_and_footers(document) -> List[str]:
    lines: List[str] = []
    try:
        for section in document.sections:
            for container in (section.header, section.first_page_header, section.even_page_header):
                for paragraph in getattr(container, "paragraphs", []):
                    text = _normalize_docx_line(paragraph.text)
                    if text:
                        lines.append(text)
                for table in getattr(container, "tables", []):
                    lines.extend(_extract_docx_table_lines(table))
            for container in (section.footer, section.first_page_footer, section.even_page_footer):
                for paragraph in getattr(container, "paragraphs", []):
                    text = _normalize_docx_line(paragraph.text)
                    if text:
                        lines.append(text)
                for table in getattr(container, "tables", []):
                    lines.extend(_extract_docx_table_lines(table))
    except Exception as exc:
        logger.warning("DOCX header/footer extraction failed: %s", exc)
    return lines


def _extract_docx_textboxes(document) -> List[str]:
    lines: List[str] = []
    try:
        xml_parts = [document.part.element.xml]
        for relation in getattr(document.part, "rels", {}).values():
            target_part = getattr(relation, "target_part", None)
            if target_part is None:
                continue
            element = getattr(target_part, "element", None)
            if element is not None:
                xml_parts.append(element.xml)
        for xml in xml_parts:
            for match in TEXTBOX_TEXT_PATTERN.findall(xml):
                text = _normalize_docx_line(match)
                if text:
                    lines.append(text)
    except Exception as exc:
        logger.warning("DOCX textbox extraction failed: %s", exc)
    return _dedupe_preserve_order(lines)


def extract_docx_tables(file_path: str) -> List[str]:
    """Extract DOCX table text row-by-row, including nested tables."""
    table_lines: List[str] = []
    try:
        import docx

        document = docx.Document(file_path)
        for table in document.tables:
            table_lines.extend(_extract_docx_table_lines(table))
    except Exception as exc:
        logger.warning("python-docx table extraction failed for %s: %s", file_path, exc)
    return _dedupe_preserve_order(table_lines)


def extract_docx(file_path: str) -> Dict[str, Any]:
    """
    Extract DOCX content using:
    1. docx2python
    2. python-docx
    3. mammoth

    Includes paragraphs, tables, headers, footers, text boxes, and multi-column content.
    """
    parser_traces: List[str] = []
    docx2python_lines: List[str] = []
    python_docx_lines: List[str] = []
    mammoth_lines: List[str] = []
    textbox_lines: List[str] = []
    header_footer_lines: List[str] = []
    merged_lines: List[str] = []

    if docx2python is not None:
        try:
            with docx2python(file_path) as extracted:
                parser_traces.append("docx2python")
                raw_text = getattr(extracted, "text", "") or ""
                docx2python_lines.extend(
                    [line for line in (_normalize_docx_line(part) for part in raw_text.splitlines()) if line]
                )
                body = getattr(extracted, "body", None)
                if body is not None:
                    docx2python_lines.extend(_flatten_docx2python_node(body))
                for container_name in ("header", "footer", "footnotes", "endnotes"):
                    container = getattr(extracted, container_name, None)
                    if container is not None:
                        docx2python_lines.extend(_flatten_docx2python_node(container))
        except Exception as exc:
            logger.warning("docx2python extraction failed for %s: %s", file_path, exc)

    try:
        import docx

        document = docx.Document(file_path)
        parser_traces.append("python-docx")
        for block in _iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                text = _normalize_docx_line(block.text)
                if text:
                    python_docx_lines.append(text)
            elif isinstance(block, Table):
                python_docx_lines.extend(_extract_docx_table_lines(block))
        header_footer_lines.extend(_extract_docx_headers_and_footers(document))
        textbox_lines.extend(_extract_docx_textboxes(document))
    except Exception as exc:
        logger.warning("python-docx block extraction failed for %s: %s", file_path, exc)

    if not docx2python_lines and not python_docx_lines and mammoth is not None:
        try:
            with open(file_path, "rb") as handle:
                mammoth_result = mammoth.extract_raw_text(handle)
            parser_traces.append("mammoth")
            mammoth_lines.extend(
                [line for line in (_normalize_docx_line(part) for part in str(mammoth_result.value or "").splitlines()) if line]
            )
        except Exception as exc:
            logger.warning("mammoth extraction failed for %s: %s", file_path, exc)

    merged_lines.extend(docx2python_lines)
    merged_lines.extend(python_docx_lines)
    merged_lines.extend(header_footer_lines)
    merged_lines.extend(textbox_lines)
    table_lines = extract_docx_tables(file_path)
    merged_lines.extend(table_lines)
    merged_lines.extend(mammoth_lines)

    merged_lines = _dedupe_preserve_order(merged_lines)
    return {
        "text": "\n".join(merged_lines),
        "paragraphs": _dedupe_preserve_order([*docx2python_lines, *python_docx_lines, *mammoth_lines]),
        "tables": table_lines,
        "textboxes": _dedupe_preserve_order(textbox_lines),
        "headers_footers": _dedupe_preserve_order(header_footer_lines),
        "parsers_used": parser_traces,
    }


def extract_docx_text(file_path: str) -> Dict[str, Any]:
    return extract_docx(file_path)


def _group_pdf_words_into_lines(words: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not words:
        return []
    ordered_words = sorted(words, key=lambda item: (item["top"], item["x0"]))
    line_groups: List[List[Dict[str, Any]]] = []
    current_group: List[Dict[str, Any]] = []
    current_top: float | None = None

    for word in ordered_words:
        word_top = float(word["top"])
        if current_top is None or abs(word_top - current_top) <= PDF_LINE_TOLERANCE:
            current_group.append(word)
            if current_top is None:
                current_top = word_top
            else:
                current_top = min(current_top, word_top)
            continue
        line_groups.append(current_group)
        current_group = [word]
        current_top = word_top

    if current_group:
        line_groups.append(current_group)

    lines: List[Dict[str, Any]] = []
    for group in line_groups:
        group = sorted(group, key=lambda item: item["x0"])
        segments: List[List[Dict[str, Any]]] = []
        current_segment: List[Dict[str, Any]] = []
        previous_x1: float | None = None

        for item in group:
            x0 = float(item["x0"])
            if previous_x1 is None or x0 - previous_x1 <= PDF_SEGMENT_GAP:
                current_segment.append(item)
            else:
                if current_segment:
                    segments.append(current_segment)
                current_segment = [item]
            previous_x1 = float(item["x1"])

        if current_segment:
            segments.append(current_segment)

        for segment in segments:
            text = " ".join(str(item["text"]).strip() for item in segment if str(item["text"]).strip()).strip()
            if not text:
                continue
            lines.append(
                {
                    "text": text,
                    "x0": min(float(item["x0"]) for item in segment),
                    "x1": max(float(item["x1"]) for item in segment),
                    "top": min(float(item["top"]) for item in segment),
                }
            )
    return lines


def _detect_pdf_column_split(lines: List[Dict[str, Any]], page_width: float) -> Optional[float]:
    starts = sorted({round(line["x0"], 1) for line in lines})
    best_split: Optional[float] = None
    best_gap = 0.0

    for left_start, right_start in zip(starts, starts[1:]):
        gap = right_start - left_start
        if gap < PDF_MIN_COLUMN_GAP:
            continue
        split = left_start + gap / 2
        left_lines = [line for line in lines if line["x0"] < split]
        right_lines = [line for line in lines if line["x0"] >= split]
        if len(left_lines) < PDF_MIN_LINES_PER_COLUMN or len(right_lines) < PDF_MIN_LINES_PER_COLUMN:
            continue
        left_avg_width = sum(line["x1"] - line["x0"] for line in left_lines) / len(left_lines)
        right_avg_width = sum(line["x1"] - line["x0"] for line in right_lines) / len(right_lines)
        if left_avg_width >= right_avg_width:
            continue
        if split <= page_width * 0.18 or split >= page_width * 0.55:
            continue
        if gap > best_gap:
            best_gap = gap
            best_split = split

    return best_split


def _extract_pdf_page_text(page: pdfplumber.page.Page) -> str:
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
    lines = _group_pdf_words_into_lines(words)
    if not lines:
        return page.extract_text() or ""

    split = _detect_pdf_column_split(lines, float(page.width))
    if split is None:
        return "\n".join(line["text"] for line in sorted(lines, key=lambda item: (item["top"], item["x0"])))

    left_lines = [line for line in lines if line["x0"] < split]
    right_lines = [line for line in lines if line["x0"] >= split]
    if not left_lines or not right_lines:
        return "\n".join(line["text"] for line in sorted(lines, key=lambda item: (item["top"], item["x0"])))

    left_first_top = min(line["top"] for line in left_lines)
    header_cutoff = max(0.0, left_first_top - 35.0)
    header_lines = [line for line in right_lines if line["top"] < header_cutoff]
    remaining_right_lines = [line for line in right_lines if line["top"] >= header_cutoff]

    ordered_lines = (
        sorted(header_lines, key=lambda item: (item["top"], item["x0"]))
        + sorted(left_lines, key=lambda item: (item["top"], item["x0"]))
        + sorted(remaining_right_lines, key=lambda item: (item["top"], item["x0"]))
    )
    return "\n".join(line["text"] for line in ordered_lines)


def _extract_pdf_text_with_pymupdf(file_path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    if fitz is None:
        return [], []

    text_parts: List[str] = []
    page_metrics: List[Dict[str, Any]] = []
    document = None
    try:
        document = fitz.open(file_path)
        for page in document:
            page_text = (page.get_text("text") or "").strip()
            if page_text:
                text_parts.append(page_text)
            blocks = page.get_text("blocks") or []
            text_blocks = [block for block in blocks if len(block) >= 5 and str(block[4]).strip()]
            upper_header_blocks = [
                block
                for block in text_blocks
                if float(block[1]) <= float(page.rect.height) * 0.22
            ]
            image_count = len(page.get_images(full=True))
            block_centers = sorted(
                float((block[0] + block[2]) / 2.0)
                for block in text_blocks
            )
            has_multi_column = False
            if len(block_centers) >= 8:
                left = [center for center in block_centers if center < float(page.rect.width) * 0.45]
                right = [center for center in block_centers if center > float(page.rect.width) * 0.55]
                has_multi_column = len(left) >= 4 and len(right) >= 4
            page_text_alpha = sum(1 for char in page_text if char.isalpha())
            is_scanned_pdf = image_count > 0 and page_text_alpha < OCR_MIN_ALPHA_CHARS
            page_metrics.append(
                {
                    "width": float(page.rect.width),
                    "height": float(page.rect.height),
                    "has_multi_column": has_multi_column,
                    "has_header_block": bool(upper_header_blocks),
                    "image_count": image_count,
                    "is_scanned_pdf": is_scanned_pdf,
                    "table_count": 0,
                    "has_table_like_structure": False,
                }
            )
    except Exception as exc:
        logger.warning("PyMuPDF extraction failed for %s: %s", file_path, exc)
    finally:
        if document is not None:
            document.close()
    return text_parts, page_metrics


def _extract_pdf_text_with_pdfplumber(file_path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    text_parts: List[str] = []
    page_metrics: List[Dict[str, Any]] = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = _extract_pdf_page_text(page) or ""
                if page_text.strip():
                    text_parts.append(page_text)
                words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
                lines = _group_pdf_words_into_lines(words)
                split = _detect_pdf_column_split(lines, float(page.width)) if lines else None
                extracted_tables = page.extract_tables() or []
                page_metrics.append(
                    {
                        "width": float(page.width),
                        "height": float(page.height),
                        "has_multi_column": split is not None,
                        "has_header_block": bool(
                            [line for line in lines if float(line.get("top") or 0.0) <= float(page.height) * 0.22]
                        ),
                        "image_count": 0,
                        "is_scanned_pdf": not page_text.strip(),
                        "table_count": len(extracted_tables),
                        "has_table_like_structure": len(extracted_tables) > 0,
                    }
                )
    except Exception as exc:
        logger.warning("pdfplumber extraction failed for %s: %s", file_path, exc)
    return text_parts, page_metrics


def detect_scanned_pdf(file_path: str) -> bool:
    if os.path.splitext(file_path)[1].lower() != ".pdf":
        return False
    text_parts, page_metrics = _extract_pdf_text_with_pymupdf(file_path)
    if any(bool(page.get("is_scanned_pdf")) for page in page_metrics):
        return True
    return not _has_meaningful_text(text_parts)


def detect_image_resume(file_path: str) -> bool:
    return os.path.splitext(file_path)[1].lower() in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


def run_ocr(file_path: str) -> List[str]:
    if detect_image_resume(file_path):
        return _extract_image_text_via_ocr(file_path)
    if os.path.splitext(file_path)[1].lower() == ".pdf":
        return _extract_pdf_text_via_ocr(file_path)
    return []


def _configure_tesseract() -> bool:
    if pytesseract is None:
        return False
    configured_cmd = os.getenv("TESSERACT_CMD", "").strip()
    if configured_cmd:
        pytesseract.pytesseract.tesseract_cmd = configured_cmd
        return os.path.exists(configured_cmd) or bool(shutil.which(configured_cmd))
    return bool(shutil.which("tesseract"))


def _is_ocr_ready() -> bool:
    return pytesseract is not None and _configure_tesseract()


def _has_meaningful_text(text_parts: List[str]) -> bool:
    combined = "\n".join(part.strip() for part in text_parts if part and part.strip())
    if not combined:
        return False
    alpha_count = sum(1 for char in combined if char.isalpha())
    if len(combined) < OCR_MIN_TEXT_LENGTH or alpha_count < OCR_MIN_ALPHA_CHARS:
        return False
    return (alpha_count / max(len(combined), 1)) >= OCR_MIN_ALPHA_RATIO


def _score_text_quality(text_parts: List[str]) -> float:
    combined = "\n".join(part.strip() for part in text_parts if part and part.strip())
    if not combined:
        return 0.0

    score = 0.0
    alpha_count = sum(1 for char in combined if char.isalpha())
    length = len(combined)
    if length >= OCR_MIN_TEXT_LENGTH:
        score += 0.25
    if alpha_count >= OCR_MIN_ALPHA_CHARS:
        score += 0.2
    score += min(alpha_count / max(length, 1), 1.0) * 0.2
    if EMAIL_PATTERN.search(combined):
        score += 0.15
    if re.search(r"(?i)\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec|\d{4})\b", combined):
        score += 0.1
    broken_penalty = 0.0
    lines = [line.strip() for line in combined.splitlines() if line.strip()]
    short_line_count = sum(1 for line in lines if len(line.split()) == 1)
    if lines and (short_line_count / len(lines)) >= 0.3:
        broken_penalty += 0.14
    early_lines = lines[:4]
    if early_lines:
        seen_contact = False
        for line in early_lines:
            if EMAIL_PATTERN.search(line) or PHONE_LINE_PATTERN.search(line):
                seen_contact = True
            elif SECTION_START_PATTERN.match(line) and not seen_contact:
                broken_penalty += 0.12
                break
    broken_penalty += min(len(BROKEN_MONTH_PATTERN.findall(combined)) * 0.08, 0.16)
    broken_penalty += min(len(SPLIT_EMAIL_ARTIFACT_PATTERN.findall(combined)) * 0.12, 0.24)
    broken_penalty += min(len(BROKEN_TOKEN_PATTERN.findall(combined[:4000])) * 0.005, 0.2)
    line_count = len(lines)
    if line_count >= 12:
        score += 0.05
    return round(max(score - broken_penalty, 0.0), 3)


def _select_best_pdf_text(
    parser_outputs: Dict[str, Dict[str, Any]]
) -> Tuple[List[str], str, Dict[str, float], Dict[str, Any]]:
    best_name = ""
    best_parts: List[str] = []
    best_score = -1.0
    scores: Dict[str, float] = {}
    best_layout: Dict[str, Any] = {}

    for parser_name in ("pymupdf", "pdfplumber"):
        parser_payload = parser_outputs.get(parser_name) or {}
        text_parts = parser_payload.get("text_parts") or []
        score = _score_text_quality(text_parts)
        scores[parser_name] = score
        layout = parser_payload.get("layout") or {}
        if score > best_score:
            best_name = parser_name
            best_parts = text_parts
            best_layout = layout
            best_score = score
            continue
        if (
            parser_name == "pdfplumber"
            and text_parts
            and layout.get("is_multi_column")
            and score >= best_score - 0.04
        ):
            best_name = parser_name
            best_parts = text_parts
            best_layout = layout
            best_score = score
            continue
        if score == best_score and text_parts:
            current_preference = PDF_PARSER_PREFERENCE.get(parser_name, 0)
            best_preference = PDF_PARSER_PREFERENCE.get(best_name, 0)
            if current_preference > best_preference:
                best_name = parser_name
                best_parts = text_parts
                best_layout = layout

    if best_parts:
        return best_parts, best_name, scores, best_layout

    for parser_name in ("pymupdf", "pdfplumber"):
        parser_payload = parser_outputs.get(parser_name) or {}
        text_parts = parser_payload.get("text_parts") or []
        if text_parts:
            return text_parts, parser_name, scores, parser_payload.get("layout") or {}

    return [], "", scores, {}


def _prepare_ocr_image(image: Image.Image) -> Image.Image:
    prepared = image.convert("L")
    return prepared.point(lambda pixel: 255 if pixel > 180 else 0)


def _render_pdf_pages_for_ocr(file_path: str) -> List[Image.Image]:
    rendered_pages: List[Image.Image] = []

    if fitz is not None:
        document = None
        try:
            document = fitz.open(file_path)
            matrix = fitz.Matrix(2, 2)
            for page in document:
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                rendered_pages.append(image)
        except Exception as exc:
            logger.warning("PyMuPDF OCR rendering failed for %s: %s", file_path, exc)
        finally:
            if document is not None:
                document.close()

    if rendered_pages:
        return rendered_pages

    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                try:
                    rendered_pages.append(page.to_image(resolution=200).original)
                except Exception as exc:
                    logger.warning("OCR page rendering failed for %s: %s", file_path, exc)
    except Exception as exc:
        logger.warning("OCR PDF open failed for %s: %s", file_path, exc)

    return rendered_pages


def _extract_pdf_text_via_ocr(file_path: str) -> List[str]:
    if not _is_ocr_ready():
        return []

    ocr_parts: List[str] = []
    try:
        for page_image in _render_pdf_pages_for_ocr(file_path):
            try:
                prepared_image = _prepare_ocr_image(page_image)
                page_text = pytesseract.image_to_string(prepared_image) or ""
                if page_text.strip():
                    ocr_parts.append(page_text)
            except Exception as exc:
                logger.warning("OCR page extraction failed for %s: %s", file_path, exc)
    except Exception as exc:
        logger.warning("OCR extraction failed for %s: %s", file_path, exc)
    return ocr_parts


def _extract_image_text_via_ocr(file_path: str) -> List[str]:
    if not _is_ocr_ready():
        return []

    try:
        with Image.open(file_path) as image:
            prepared_image = _prepare_ocr_image(image)
            image_text = pytesseract.image_to_string(prepared_image) or ""
            if image_text.strip():
                return [image_text]
    except Exception as exc:
        logger.warning("Image OCR extraction failed for %s: %s", file_path, exc)
    return []


def _extract_pdf_text_with_pypdf(file_path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    # Compatibility shim kept intentionally unused in the production extraction stack.
    return [], []


def _default_layout_signals() -> Dict[str, Any]:
    return infer_layout_signals(text_parts=[], page_metrics=[])


def _record_stage_time(performance: Dict[str, float], stage: str, started_at: float) -> None:
    performance[stage] = round((time.perf_counter() - started_at) * 1000.0, 2)


def get_parser_runtime_status() -> Dict[str, Any]:
    layout_runtime = get_layout_runtime_status()
    return {
        "pymupdf_available": fitz is not None,
        "pdfplumber_available": True,
        "mammoth_available": mammoth is not None,
        "pytesseract_available": pytesseract is not None,
        "ocr_ready": _is_ocr_ready(),
        "tesseract_cmd": getattr(getattr(pytesseract, "pytesseract", None), "tesseract_cmd", "") if pytesseract is not None else "",
        **layout_runtime,
    }


def extract_document(file_path: str) -> Dict[str, Any]:
    if not file_path or not os.path.exists(file_path):
        return {"text": "", "layout": _default_layout_signals(), "tables": [], "metadata": {}}

    file_ext = os.path.splitext(file_path)[1].lower()
    text_parts: List[str] = []
    layout_signals = _default_layout_signals()
    performance: Dict[str, float] = {
        "file_upload_ms": 0.0,
        "pdf_extraction_ms": 0.0,
        "docx_extraction_ms": 0.0,
        "ocr_ms": 0.0,
    }

    try:
        if file_ext == ".pdf":
            parser_outputs: Dict[str, Dict[str, Any]] = {}
            primary_started_at = time.perf_counter()
            primary_text_parts, primary_page_metrics = _extract_pdf_text_with_pymupdf(file_path)
            _record_stage_time(performance, "pdf_extraction_ms", primary_started_at)
            parser_outputs["pymupdf"] = {
                "text_parts": primary_text_parts,
                "page_metrics": primary_page_metrics,
                "layout": infer_layout_signals(text_parts=primary_text_parts, page_metrics=primary_page_metrics),
            }
            text_parts, selected_parser, parser_scores, selected_layout = _select_best_pdf_text(parser_outputs)
            layout_signals = selected_layout or _default_layout_signals()
            native_score = _score_text_quality(text_parts)
            scanned_pdf = bool(layout_signals.get("is_scanned_pdf")) or detect_scanned_pdf(file_path)

            should_run_pdf_fallback = (
                not _has_meaningful_text(text_parts)
                or native_score < 0.45
                or bool(layout_signals.get("is_multi_column"))
            )
            if should_run_pdf_fallback:
                fallback_started_at = time.perf_counter()
                fallback_text_parts, fallback_page_metrics = _extract_pdf_text_with_pdfplumber(file_path)
                performance["pdf_extraction_ms"] = round(
                    performance.get("pdf_extraction_ms", 0.0) + (time.perf_counter() - fallback_started_at) * 1000.0,
                    2,
                )
                parser_outputs["pdfplumber"] = {
                    "text_parts": fallback_text_parts,
                    "page_metrics": fallback_page_metrics,
                    "layout": infer_layout_signals(text_parts=fallback_text_parts, page_metrics=fallback_page_metrics),
                }
                text_parts, selected_parser, parser_scores, selected_layout = _select_best_pdf_text(parser_outputs)
                layout_signals = selected_layout or layout_signals
                native_score = _score_text_quality(text_parts)
                scanned_pdf = bool(layout_signals.get("is_scanned_pdf")) or scanned_pdf

            if scanned_pdf or not _has_meaningful_text(text_parts):
                ocr_started_at = time.perf_counter()
                ocr_parts = run_ocr(file_path)
                _record_stage_time(performance, "ocr_ms", ocr_started_at)
                if _has_meaningful_text(ocr_parts):
                    ocr_score = _score_text_quality(ocr_parts)
                    if ocr_score > native_score + 0.08:
                        text_parts = ocr_parts
                        layout_signals["ocr_applied"] = True

            if selected_parser:
                logger.info(
                    "Selected PDF parser '%s' for %s with scores: %s and layout: %s",
                    selected_parser,
                    file_path,
                    parser_scores,
                    layout_signals,
                )
            layout_signals["is_scanned_pdf"] = bool(layout_signals.get("is_scanned_pdf")) or scanned_pdf

        elif file_ext == ".docx":
            try:
                docx_started_at = time.perf_counter()
                docx_payload = extract_docx(file_path)
                _record_stage_time(performance, "docx_extraction_ms", docx_started_at)
                extracted_text = str(docx_payload.get("text") or "").strip()
                if extracted_text:
                    text_parts.append(extracted_text)
                layout_signals = infer_layout_signals(text_parts=text_parts, page_metrics=[])
                return {
                    "text": "\n".join(part.strip() for part in text_parts if part and part.strip()),
                    "layout": layout_signals,
                    "tables": docx_payload.get("tables") or [],
                    "metadata": {
                        "format": "docx",
                        "parsers_used": docx_payload.get("parsers_used") or [],
                        "paragraph_count": len(docx_payload.get("paragraphs") or []),
                        "table_row_count": len(docx_payload.get("tables") or []),
                        "textbox_count": len(docx_payload.get("textboxes") or []),
                        "header_footer_count": len(docx_payload.get("headers_footers") or []),
                        "performance": performance,
                    },
                }
            except Exception as exc:
                logger.warning("python-docx extraction failed for %s: %s", file_path, exc)

        elif file_ext == ".doc":
            try:
                import subprocess

                result = subprocess.run(["antiword", file_path], capture_output=True, text=True)
                if result.returncode == 0 and result.stdout.strip():
                    text_parts.append(result.stdout)
                layout_signals = infer_layout_signals(text_parts=text_parts, page_metrics=[])
            except Exception as exc:
                logger.warning("antiword extraction failed for %s: %s", file_path, exc)

        elif file_ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
            ocr_started_at = time.perf_counter()
            text_parts = run_ocr(file_path)
            _record_stage_time(performance, "ocr_ms", ocr_started_at)
            layout_signals = infer_layout_signals(text_parts=text_parts, page_metrics=[])
            layout_signals["is_scanned_pdf"] = False
            layout_signals["ocr_applied"] = True

        if not text_parts:
            with open(file_path, "rb") as handle:
                binary_text = handle.read().decode("utf-8", errors="ignore")
                if binary_text.strip():
                    text_parts.append(binary_text)
            layout_signals = infer_layout_signals(text_parts=text_parts, page_metrics=[])
    except Exception as exc:
        logger.exception("Resume text extraction failed for %s: %s", file_path, exc)
        return {"text": "", "layout": _default_layout_signals()}

    return {
        "text": "\n".join(part.strip() for part in text_parts if part and part.strip()),
        "layout": layout_signals,
        "tables": [],
        "metadata": {
            "format": file_ext.lstrip(".") or "text",
            "ocr_applied": bool(layout_signals.get("ocr_applied")),
            "is_image_resume": detect_image_resume(file_path),
            "performance": performance,
        },
    }


def extract_text(file_path: str) -> str:
    return str(extract_document(file_path).get("text") or "")


def _normalize_name_candidate(value: str) -> str:
    candidate = re.sub(r"\s+", " ", (value or "").strip(" ,.-"))
    if not candidate:
        return ""
    words = candidate.split()
    if not (2 <= len(words) <= 4):
        return ""
    if any(any(char.isdigit() for char in word) for word in words):
        return ""
    lowered_words = [word.lower().strip(".,") for word in words]
    if any(word in INVALID_NAME_TOKENS for word in lowered_words):
        return ""
    if NAME_COMPANY_PATTERN.search(candidate):
        return ""
    if extract_normalized_location(candidate, use_spacy=False) or validate_location(candidate):
        return ""
    if NAME_CONTEXT_ROLE_PATTERN.search(candidate):
        return ""
    if "@" in candidate or PHONE_LINE_PATTERN.search(candidate):
        return ""
    if not all(word.replace(".", "").replace("'", "").isalpha() for word in words):
        return ""
    if not all(word.isupper() or word[:1].isupper() for word in words):
        return ""
    return " ".join(word if len(word) == 1 else word.title() for word in words)


def _extract_inline_header_name(line: str) -> str:
    candidate_line = re.sub(r"\s+", " ", (line or "").strip())
    if not candidate_line:
        return ""

    candidate_line = re.sub(r"(?i)^contact\s+", "", candidate_line).strip()
    candidate_line = INLINE_CONTACT_PATTERN.split(candidate_line, maxsplit=1)[0].strip(" ,|-")

    tokens = [token.strip(" ,.-") for token in candidate_line.split() if token.strip(" ,.-")]
    collected: List[str] = []
    for token in tokens:
        lowered = token.lower()
        normalized_candidate = _normalize_name_candidate(" ".join(collected + [token]))
        if lowered in NAME_STOP_TOKENS:
            break
        if any(char.isdigit() for char in token) or "@" in token:
            break
        if lowered in INVALID_NAME_TOKENS:
            break
        if token.isupper() or token[:1].isupper():
            collected.append(token)
        else:
            break
        if len(collected) >= 4:
            break
        if normalized_candidate and len(collected) >= 2:
            next_index = len(collected)
            if next_index < len(tokens) and tokens[next_index].lower() in NAME_STOP_TOKENS:
                break

    leading_name = _normalize_name_candidate(" ".join(collected))
    if leading_name:
        return leading_name

    for segment in HEADER_NAME_SPLIT_PATTERN.split(candidate_line):
        normalized = _normalize_name_candidate(segment)
        if normalized:
            return normalized

    return ""


def _header_name_candidates(text: str) -> List[str]:
    return _extract_contact_zone_lines(text)


def _extract_contact_zone_lines(text: str) -> List[str]:
    cleaned_text = clean_text_pipeline(text or "")
    sections = segment_resume_sections(cleaned_text)
    structural_text = normalize_document_structure(text or "")
    structural_sections = segment_resume_sections(structural_text)

    structural_header_lines = [
        line.strip() for line in structural_sections.get("header", "").splitlines() if line.strip()
    ]
    header_lines = [line.strip() for line in sections.get("header", "").splitlines() if line.strip()]
    if structural_header_lines or header_lines:
        return structural_header_lines or header_lines

    raw_lines = [line.strip() for line in structural_text.splitlines() if line.strip()]
    contact_lines: List[str] = []
    for line in raw_lines[:10]:
        if SECTION_START_PATTERN.match(line):
            break
        contact_lines.append(line)
    return contact_lines or raw_lines[:5]


def _extract_contact_zone_text(text: str) -> str:
    return "\n".join(_extract_contact_zone_lines(text))


def _extract_phone(text: str) -> str:
    source_text = _extract_contact_zone_text(text)
    for pattern in PHONE_PATTERNS:
        matches = re.findall(pattern, source_text or "")
        if matches:
            return matches[0].strip()
    return ""


def _extract_email(text: str) -> str:
    source_text = _extract_contact_zone_text(text)
    email = extract_normalized_email(source_text or "")
    if email:
        return email
    structured_text = normalize_document_structure(text or "")
    early_lines = [line.strip() for line in structured_text.splitlines() if line.strip()][:20]
    bounded_early_lines: List[str] = []
    for line in early_lines:
        if bounded_early_lines and SECTION_START_PATTERN.match(line):
            break
        bounded_early_lines.append(line)
    labeled_candidates: List[str] = []
    for line in bounded_early_lines:
        if SECTION_START_PATTERN.match(line):
            break
        if re.match(r"(?i)^(?:email|mail|e-mail)\s*[:\-]", line):
            labeled_candidates.append(line)
    if labeled_candidates:
        labeled_email = extract_normalized_email("\n".join(labeled_candidates))
        if labeled_email:
            return labeled_email
    return extract_normalized_email("\n".join(bounded_early_lines)) or ""


def _extract_name_with_spacy(text: str) -> str:
    if not SPACY_AVAILABLE:
        return ""
    search_zones = [
        "\n".join(_header_name_candidates(text)[:8]),
        "\n".join(line.strip() for line in normalize_document_structure(text or "").splitlines()[:20] if line.strip()),
    ]
    for zone in search_zones:
        if not zone.strip():
            continue
        doc = get_section_doc(zone)
        if doc is None:
            continue
        for ent in doc.ents:
            if ent.label_ != "PERSON":
                continue
            candidate = _normalize_name_candidate(ent.text)
            if candidate:
                return candidate
    return ""


def _extract_name(text: str, original_filename: Optional[str] = None) -> str:
    for line in _header_name_candidates(text)[:2]:
        lowered_line = line.strip().lower()
        if lowered_line in {"contact details", "contact information"}:
            continue
        contact_match = re.search(r"(?i)\bcontact\s+(?P<name>[A-Z][A-Za-z'`.-]+(?:\s+[A-Z][A-Za-z'`.-]+){1,3})\b", line)
        if contact_match:
            contact_name = _normalize_name_candidate(contact_match.group("name"))
            if contact_name:
                return contact_name
        if SECTION_START_PATTERN.match(line):
            continue
        label_match = NAME_LABEL_PATTERN.match(line)
        if label_match:
            labeled_name = _normalize_name_candidate(label_match.group("value"))
            if labeled_name:
                return labeled_name
        prefix_segment = re.split(r"\s+\|\s+|\s+[•·]\s+", line, maxsplit=1)[0].strip()
        normalized = _normalize_name_candidate(prefix_segment)
        if normalized:
            return normalized
        inline_prefix_name = _extract_inline_header_name(prefix_segment)
        if inline_prefix_name:
            return inline_prefix_name
        normalized = _normalize_name_candidate(line)
        if normalized:
            return normalized
        inline_header_name = _extract_inline_header_name(line)
        if inline_header_name:
            return inline_header_name
    normalized_source = normalize_document_structure(text or "")
    extracted_name = extract_name(normalized_source)
    if extracted_name:
        return extracted_name
    spacy_name = _extract_name_with_spacy(text)
    if spacy_name:
        return spacy_name
    raw_lines = [line.strip() for line in normalize_document_structure(text or "").splitlines() if line.strip()]
    for index, line in enumerate(raw_lines[:80]):
        lowered_line = line.strip().lower()
        if lowered_line in {"contact details", "contact information"}:
            continue
        if SECTION_START_PATTERN.match(line):
            continue
        if EMAIL_PATTERN.search(line) or PHONE_LINE_PATTERN.search(line):
            continue
        normalized = _normalize_name_candidate(line)
        if not normalized:
            continue
        next_line = raw_lines[index + 1].strip() if index + 1 < len(raw_lines) else ""
        if next_line and NAME_CONTEXT_ROLE_PATTERN.search(next_line):
            return normalized
    for line in raw_lines[:60]:
        lowered_line = line.strip().lower()
        if lowered_line in {"contact details", "contact information"}:
            continue
        if SECTION_START_PATTERN.match(line):
            continue
        if EMAIL_PATTERN.search(line) or PHONE_LINE_PATTERN.search(line):
            continue
        label_match = NAME_LABEL_PATTERN.match(line)
        if label_match:
            labeled_name = _normalize_name_candidate(label_match.group("value"))
            if labeled_name:
                return labeled_name
        normalized = _normalize_name_candidate(line)
        if normalized:
            return normalized
        inline_header_name = _extract_inline_header_name(line)
        if inline_header_name:
            return inline_header_name
    return ""


def _score_name_confidence(name: str) -> float:
    normalized = _normalize_name_candidate(name)
    if not normalized:
        return 0.0
    score = 0.45
    if 2 <= len(normalized.split()) <= 4:
        score += 0.2
    if not any(token.lower() in INVALID_NAME_TOKENS for token in normalized.split()):
        score += 0.15
    if not NAME_COMPANY_PATTERN.search(normalized) and not NAME_CONTEXT_ROLE_PATTERN.search(normalized):
        score += 0.2
    return round(min(score, 1.0), 2)


def _score_email_confidence(email: str) -> float:
    return 1.0 if _extract_email(email) else 0.0


def _score_phone_confidence(phone: str) -> float:
    return 1.0 if _extract_phone(phone) else 0.0


def _score_skill_confidence(skills: List[str], entities: Dict[str, Any], from_explicit_section: bool) -> float:
    if not skills:
        return 0.0
    score = 0.35
    if from_explicit_section:
        score += 0.35
    entity_skills = {str(skill).strip().lower() for skill in entities.get("skills", []) if str(skill).strip()}
    matched = sum(1 for skill in skills if str(skill).strip().lower() in entity_skills)
    if matched:
        score += min(0.2, matched * 0.05)
    if len(skills) >= 3:
        score += 0.1
    return round(min(score, 1.0), 2)


def _confidence_to_percent(value: float) -> int:
    return int(round(max(0.0, min(1.0, value)) * 100))


def _score_location_confidence(location: str) -> float:
    if not location:
        return 0.0
    score = 0.45
    if "," in location:
        score += 0.2
    if not any(char.isdigit() for char in location):
        score += 0.15
    if 1 <= len(location.split()) <= 3:
        score += 0.2
    return round(min(score, 1.0), 2)


def _rerun_low_confidence_fields(
    result: Dict[str, Any],
    *,
    entities: Dict[str, Any],
    normalized_text: str,
    original_filename: Optional[str],
) -> Dict[str, Any]:
    updated = dict(result)
    field_confidence = dict(updated.get("field_confidence") or {})

    if field_confidence.get("name", 0.0) < CONFIDENCE_RETRY_THRESHOLD:
        updated["name"] = ""

    if field_confidence.get("location", 0.0) < CONFIDENCE_RETRY_THRESHOLD:
        location_candidates = [
            updated.get("sections", {}).get("header", ""),
            entities.get("top_location", ""),
            normalized_text,
        ]
        for candidate in location_candidates:
            extracted_location = extract_normalized_location(candidate, use_spacy=False)
            if not extracted_location and SPACY_AVAILABLE:
                extracted_location = extract_normalized_location(candidate, use_spacy=True)
            if not extracted_location:
                for line in [line.strip() for line in str(candidate or "").splitlines() if line.strip()][:8]:
                    extracted_location = validate_location(line)
                    if extracted_location:
                        break
            if extracted_location:
                updated["location"] = extracted_location
                break

    if field_confidence.get("skills", 0.0) < CONFIDENCE_RETRY_THRESHOLD:
        fallback_skill_candidates: List[str] = []
        normalized_corpus = normalized_text.lower()
        for skill in entities.get("skills", []) or []:
            normalized_skill = str(skill).strip().lower()
            if not normalized_skill:
                continue
            if len(normalized_skill.split()) > 3:
                continue
            if normalized_skill not in normalized_corpus:
                continue
            fallback_skill_candidates.append(skill)
        merged_skills = dedupe_strings([*(updated.get("skills") or []), *fallback_skill_candidates])
        if merged_skills:
            updated["skills"] = merged_skills

    return updated


def extract_skills(text: str) -> List[str]:
    info = extract_resume_information(text)
    return info.get("skills", [])


def extract_experience(text: str) -> List[Dict[str, Any]]:
    info = extract_resume_information(text)
    return info.get("experience", [])


def calculate_total_experience(experience_entries: List[Dict[str, Any]]) -> float:
    if not experience_entries:
        return 0.0
    ranges = []
    for entry in experience_entries:
        start = parse_date(entry.get("start_date", ""), is_end=False)
        end = parse_date(entry.get("end_date", ""), is_end=True)
        if not start or not end or end < start:
            continue
        ranges.append((start, end))
    if not ranges:
        return 0.0
    return compute_total_experience(ranges)


def extract_current_company(experience_entries: List[Dict[str, Any]]) -> str:
    return experience_entries[0].get("company") if experience_entries else ""


def extract_location(text: str) -> str:
    info = extract_resume_information(text)
    return info.get("location", "")


def _classify_experience_level(total_experience_years: Optional[float]) -> str:
    if total_experience_years is None:
        return ""
    if total_experience_years <= 2:
        return "Junior"
    if total_experience_years <= 5:
        return "Mid-level"
    if total_experience_years <= 10:
        return "Senior"
    return "Lead/Expert"


def parse_resume_text(
    raw_text: str,
    original_filename: Optional[str] = None,
    layout_signals: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    layout_signals = layout_signals or _default_layout_signals()
    cleaned_text = clean_text(raw_text) if raw_text else ""
    normalized_text = clean_text_pipeline(cleaned_text) if cleaned_text else ""
    extracted_info = extract_resume_information(normalized_text or cleaned_text) if cleaned_text else {
        "sections": segment_resume_sections(""),
        "skills": [],
        "experience": [],
        "projects": [],
        "education": [],
        "certifications": [],
        "location": "",
        "current_company": None,
        "current_role": None,
        "designation": None,
        "experience_years": None,
        "total_experience_years": None,
        "experience_level": "",
        "languages": [],
    }

    sections = extracted_info.get("sections", {}) or segment_resume_sections(normalized_text)
    resume_type_payload = detect_resume_type("\n".join(filter(None, [sections.get("summary", ""), sections.get("skills", ""), sections.get("experience", ""), normalized_text[:2000]])))
    entities = extract_resume_entities(
        normalized_text,
        header_text=sections.get("header", ""),
        skills_text=sections.get("skills", ""),
        experience_text=sections.get("experience", ""),
    )
    extracted_name = _extract_name(normalized_text or raw_text, original_filename)
    contact_email = _extract_email(cleaned_text or raw_text)
    contact_phone = _extract_phone(cleaned_text or raw_text)
    merged_skills = dedupe_strings(extracted_info.get("skills", []) or [])

    result = {
        "name": extracted_name or None,
        "email": contact_email,
        "phone": contact_phone,
        "skills": merged_skills,
        "education": extracted_info.get("education", []),
        "experience": extracted_info.get("experience", []),
        "projects": extracted_info.get("projects", []),
        "certifications": extracted_info.get("certifications", []),
        "total_experience": extracted_info.get("total_experience"),
        "total_experience_years": extracted_info.get("total_experience_years"),
        "total_experience_months": extracted_info.get("total_experience_months"),
        "experience_years": extracted_info.get("experience_years"),
        "experience_level": extracted_info.get("experience_level") or _classify_experience_level(extracted_info.get("total_experience_years")),
        "current_company": extracted_info.get("current_company"),
        "current_role": extracted_info.get("current_role"),
        "designation": extracted_info.get("designation"),
        "location": extracted_info.get("location", ""),
        "languages": extracted_info.get("languages", []),
        "summary": sections.get("summary", ""),
        "sections": sections,
        "experience_entries": extracted_info.get("experience", []),
        "entities": entities,
        "resume_type": resume_type_payload.get("resume_type", "general"),
        "resume_type_scores": resume_type_payload.get("resume_type_scores", {}),
        "personal_details": {
            "name": extracted_name or None,
            "email": contact_email,
            "phone": contact_phone,
            "location": extracted_info.get("location", ""),
            "current_company": extracted_info.get("current_company"),
            "current_role": extracted_info.get("current_role"),
        },
        "pipeline_layers": {
            "file_upload": True,
            "format_detection": True,
            "text_extraction": bool(raw_text.strip()),
            "ocr_layer": bool(layout_signals.get("is_scanned_pdf")) or bool(layout_signals.get("ocr_applied")),
            "layout_detection": bool(layout_signals),
            "section_detection": bool(sections),
            "header_detection": bool(sections.get("header", "").strip()) or bool(layout_signals.get("has_header_block")),
            "entity_extraction": bool(entities.get("persons") or entities.get("skills") or entities.get("locations")),
            "field_extraction": True,
            "validation": True,
            "deduplication": True,
            "confidence_scoring": True,
            "post_processing": True,
            "final_output": True,
        },
        "layout_signals": layout_signals,
        "runtime_status": get_parser_runtime_status(),
        "full_text": cleaned_text,
        "normalized_text": normalized_text,
        "raw_text": raw_text,
        "field_confidence": {
            "name": 0.0,
            "email": 0.0,
            "phone": 0.0,
            "experience": extracted_info.get("experience_extraction_confidence", 0.0),
        },
    }
    result = apply_postprocessing(result)
    result = validate_parsed_fields(result)
    result["field_confidence"]["name"] = _score_name_confidence(result.get("name", ""))
    if result["field_confidence"]["name"] < CONFIDENCE_RETRY_THRESHOLD:
        result["name"] = None
        result["personal_details"]["name"] = None
        result["field_confidence"]["name"] = 0.0
    result["field_confidence"]["email"] = _score_email_confidence(result.get("email", ""))
    result["field_confidence"]["phone"] = _score_phone_confidence(result.get("phone", ""))
    result["field_confidence"]["location"] = _score_location_confidence(result.get("location", ""))
    result["field_confidence"]["skills"] = _score_skill_confidence(
        result.get("skills", []),
        entities,
        bool(sections.get("skills", "").strip()),
    )
    if not result.get("experience_entries") and result.get("total_experience_years") is None:
        result["field_confidence"]["experience"] = 1.0
    if any(score < CONFIDENCE_RETRY_THRESHOLD for score in result["field_confidence"].values()):
        result = _rerun_low_confidence_fields(
            result,
            entities=entities,
            normalized_text=normalized_text,
            original_filename=original_filename,
        )
        result = apply_postprocessing(result)
        result = validate_parsed_fields(result)
        result["field_confidence"]["name"] = _score_name_confidence(result.get("name", ""))
        if result["field_confidence"]["name"] < CONFIDENCE_RETRY_THRESHOLD:
            result["name"] = None
            result["personal_details"]["name"] = None
            result["field_confidence"]["name"] = 0.0
        result["field_confidence"]["email"] = _score_email_confidence(result.get("email", ""))
        result["field_confidence"]["phone"] = _score_phone_confidence(result.get("phone", ""))
        result["field_confidence"]["location"] = _score_location_confidence(result.get("location", ""))
        result["field_confidence"]["skills"] = _score_skill_confidence(
            result.get("skills", []),
            entities,
            bool(sections.get("skills", "").strip()),
        )
        if not result.get("experience_entries") and result.get("total_experience_years") is None:
            result["field_confidence"]["experience"] = 1.0
    result["confidence"] = {
        field: _confidence_to_percent(score)
        for field, score in (result.get("field_confidence") or {}).items()
    }
    result["pipeline_summary"] = {
        "format": "text",
        "retry_applied": any(
            value < _confidence_to_percent(CONFIDENCE_RETRY_THRESHOLD)
            for value in result["confidence"].values()
        ),
        "all_fields_above_80": all(value >= 80 for value in result["confidence"].values()),
    }
    logger.info(
        "Parsed resume: name=%s, skills=%s, experience_years=%s, current_company=%s, location=%s",
        result["name"],
        len(result["skills"]),
        result["total_experience_years"],
        result["current_company"],
        result["location"],
    )
    return result


def parse_resume(file_path: str, original_filename: Optional[str] = None) -> Dict[str, Any]:
    document_payload = extract_document(file_path)
    raw_text = str(document_payload.get("text") or "")
    layout_signals = document_payload.get("layout") or _default_layout_signals()
    parsed_resume = parse_resume_text(raw_text, original_filename=original_filename, layout_signals=layout_signals)
    parsed_resume["document_tables"] = document_payload.get("tables") or []
    parsed_resume["document_metadata"] = document_payload.get("metadata") or {}
    parsed_resume["pipeline_summary"]["format"] = (
        parsed_resume["document_metadata"].get("format")
        or os.path.splitext(file_path)[1].lstrip(".")
        or "text"
    )
    return parsed_resume


if __name__ == "__main__":  # pragma: no cover - example usage
    import json
    import sys

    if len(sys.argv) < 2:
        print("Usage: python -m ats.extraction.resume_parser <resume-file>")
        raise SystemExit(1)

    parsed = parse_resume(sys.argv[1], original_filename=os.path.basename(sys.argv[1]))
    print(json.dumps(parsed, indent=2, default=str))
