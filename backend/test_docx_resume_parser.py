import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ats.extraction.resume_parser import extract_docx_tables, extract_docx_text, parse_resume  # noqa: E402
from ats.preprocessing.text_cleaning import clean_text  # noqa: E402


class DocxResumeParserTests(unittest.TestCase):
    def _build_sample_docx(self) -> str:
        document = Document()
        document.add_paragraph("John Doe")
        document.add_paragraph("john.doe@example.com")

        details_table = document.add_table(rows=2, cols=2)
        details_table.cell(0, 0).text = "Phone"
        details_table.cell(0, 1).text = "+1 555 123 4567"
        details_table.cell(1, 0).text = "Location"
        details_table.cell(1, 1).text = "New York, USA"

        resume_table = document.add_table(rows=2, cols=2)
        resume_table.cell(0, 0).text = "Education"
        resume_table.cell(0, 1).text = "B.Tech | State University | 2019"
        resume_table.cell(1, 0).text = "Skills"
        nested_table = resume_table.cell(1, 1).add_table(rows=2, cols=2)
        nested_table.cell(0, 0).text = "Python"
        nested_table.cell(0, 1).text = "FastAPI"
        nested_table.cell(1, 0).text = "SQL"
        nested_table.cell(1, 1).text = "Docker"

        temp_dir = tempfile.mkdtemp()
        file_path = str(Path(temp_dir) / "sample_resume.docx")
        document.save(file_path)
        return file_path

    def test_extract_docx_tables_preserves_rows(self):
        file_path = self._build_sample_docx()

        table_lines = extract_docx_tables(file_path)

        self.assertTrue(any("Phone | +1 555 123 4567" in line for line in table_lines))
        self.assertTrue(any("Location | New York, USA" in line for line in table_lines))
        self.assertTrue(any("Education | B.Tech | State University | 2019" in line for line in table_lines))
        self.assertTrue(any("Skills | Python | FastAPI" in line or "Skills | SQL | Docker" in line for line in table_lines))

    def test_extract_docx_text_merges_paragraphs_and_tables(self):
        file_path = self._build_sample_docx()

        payload = extract_docx_text(file_path)
        text = payload["text"]

        self.assertIn("John Doe", text)
        self.assertIn("john.doe@example.com", text)
        self.assertIn("Phone | +1 555 123 4567", text)
        self.assertGreaterEqual(len(payload["parsers_used"]), 1)

    def test_parse_resume_includes_structured_sections_from_docx_tables(self):
        file_path = self._build_sample_docx()

        parsed = parse_resume(file_path, original_filename="sample_resume.docx")

        self.assertEqual(parsed["name"], "John Doe")
        self.assertEqual(parsed["email"], "john.doe@example.com")
        self.assertIn("python", parsed["skills"])
        self.assertTrue(parsed["document_tables"])
        self.assertIn("personal_details", parsed)
        self.assertIn("education", parsed)

    def test_clean_text_keeps_structure_and_colons(self):
        raw_text = "Name:\u00a0John Doe\t\tSkills:\tPython  \n\nEducation:\tB.Tech"

        cleaned = clean_text(raw_text)

        self.assertIn("Name: John Doe", cleaned)
        self.assertIn("Skills:", cleaned)
        self.assertIn("Python", cleaned)
        self.assertIn("\n\nEducation:", cleaned)


if __name__ == "__main__":
    unittest.main()
